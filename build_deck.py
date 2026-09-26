import sys
import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path):
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    C_NAVY = RGBColor(15, 23, 42)          # #0f172a
    C_NAVY_LIGHT = RGBColor(30, 41, 59)    # #1e293b
    C_WHITE = RGBColor(255, 255, 255)
    C_BG_LIGHT = RGBColor(248, 250, 252)   # #f8fafc
    C_CARD_BG = RGBColor(255, 255, 255)
    C_CARD_BORDER = RGBColor(226, 232, 240)# #e2e8f0
    C_TEXT_DARK = RGBColor(15, 23, 42)
    C_TEXT_MUTED = RGBColor(100, 116, 139) # #64748b
    C_BLUE = RGBColor(2, 132, 199)         # #0284c7
    C_BLUE_DARK = RGBColor(3, 105, 161)
    C_GREEN = RGBColor(22, 163, 74)        # #16a34a
    C_GREEN_BG = RGBColor(240, 253, 244)   # #f0fdf4
    C_AMBER = RGBColor(217, 119, 6)        # #d97706
    C_AMBER_BG = RGBColor(254, 243, 199)   # #fef3c7
    C_RED = RGBColor(220, 38, 38)          # #dc2626
    C_RED_BG = RGBColor(254, 242, 242)     # #fef2f2

    def add_bg(slide, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        return shape

    def add_header(slide, title_text, category_text, slide_num, total_slides=10):
        # Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(10.0), Inches(0.3))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10.5)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_BLUE
        p_cat.font.name = "Arial"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.74), Inches(10.5), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(23)
        p_title.font.bold = True
        p_title.font.color.rgb = C_TEXT_DARK
        p_title.font.name = "Arial"

        # Footer (Breadcrumb & Slide Number)
        ft_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
        ft_tf = ft_box.text_frame
        ft_tf.margin_left = ft_tf.margin_right = ft_tf.margin_top = ft_tf.margin_bottom = 0
        p_left = ft_tf.paragraphs[0]
        p_left.text = f"ocr-processor · In-App Automation Testing Demonstration          |          Slide {slide_num} of {total_slides}"
        p_left.font.size = Pt(9.5)
        p_left.font.color.rgb = C_TEXT_MUTED
        p_left.font.name = "Arial"

    # =========================================================================
    # SLIDE 1: TITLE SLIDE (Dark Navy)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1, C_NAVY)

    # Decorative top accent line
    top_line = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.15), Inches(1.5), Inches(0.08))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = C_BLUE
    top_line.line.fill.background()

    # Title box
    tbox = s1.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.7), Inches(4.5))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "In-App Automation Testing Platform"
    p0.font.size = Pt(40)
    p0.font.bold = True
    p0.font.color.rgb = C_WHITE
    p0.font.name = "Arial"
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = "Bridging the Developer–QA Gap with Interactive Test Execution, Live Evidence & Requirement Traceability"
    p1.font.size = Pt(19)
    p1.font.color.rgb = RGBColor(148, 163, 184) # slate-400
    p1.font.name = "Arial"
    p1.space_after = Pt(36)

    p2 = tf.add_paragraph()
    p2.text = "Project: ocr-processor (Full-Stack Document Intelligence with Tesseract OCR & Spring AI)\nDesigned, Architected & Implemented with IBM Bob"
    p2.font.size = Pt(14)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(56, 189, 248) # sky-400
    p2.font.name = "Arial"

    # Bottom pill badges
    badges = [
        ("101 Automated Cases", C_BLUE, C_NAVY_LIGHT),
        ("Triple-Evidence Model", C_GREEN, C_NAVY_LIGHT),
        ("73.1% Req Coverage", C_AMBER, C_NAVY_LIGHT),
        ("Co-Piloted with IBM Bob", C_WHITE, C_BLUE)
    ]
    for idx, (label, text_c, bg_c) in enumerate(badges):
        bx = Inches(0.8 + idx * 2.95)
        by = Inches(5.8)
        pill = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, Inches(2.75), Inches(0.65))
        pill.fill.solid()
        pill.fill.fore_color.rgb = bg_c
        pill.line.color.rgb = text_c
        pill.line.width = Pt(1.5)
        ptf = pill.text_frame
        ptf.word_wrap = True
        pp = ptf.paragraphs[0]
        pp.text = label
        pp.alignment = PP_ALIGN.CENTER
        pp.font.size = Pt(12)
        pp.font.bold = True
        pp.font.color.rgb = text_c
        pp.font.name = "Arial"

    # =========================================================================
    # SLIDE 2: THE PROBLEM (Status Quo Friction)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2, C_BG_LIGHT)
    add_header(s2, "The Core Testing Problem: Status Quo Friction", "THE PROBLEM & CHALLENGE", 2)

    cards_data_s2 = [
        ("Disconnected Feedback Loop",
         "Testing happens apart from development in separate pipelines or isolated QA sprints.",
         ["Developers push changes without immediate verification",
          "Failure reports arrive days later, breaking mental context",
          "Engineers spend entire cycles context-switching to reproduce",
          "Significant delay in overall SDLC release velocity"],
         C_RED, C_RED_BG),

        ("Ambiguous Bug Reports",
         "Bug notices routinely lack the concrete evidence needed for rapid root-cause diagnosis.",
         ["Reports say 'Case Failed' with no runtime parameters",
          "Missing test data inputs, payloads, and mock configurations",
          "No screenshots showing what the UI actually rendered",
          "Heavy back-and-forth: 'Cannot reproduce on my machine'"],
         C_AMBER, C_AMBER_BG),

        ("Late Regression Discovery",
         "New features silently compromise existing functional flows without early warning.",
         ["Unit tests only exercise isolated, mocked layers",
          "End-to-end user journeys are deferred to manual QA passes",
          "QA team becomes the first group to discover regressions",
          "Costly and stressful rework right before production release"],
         C_NAVY, RGBColor(241, 245, 249))
    ]

    for i, (title, subtitle, bullets, acc_c, bg_c) in enumerate(cards_data_s2):
        cx = Inches(0.8 + i * 4.0)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.6), Inches(3.75), Inches(5.1))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_BG
        card.line.color.rgb = C_CARD_BORDER
        card.line.width = Pt(1.5)

        abar = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + Inches(0.25), Inches(1.85), Inches(3.25), Inches(0.08))
        abar.fill.solid()
        abar.fill.fore_color.rgb = acc_c
        abar.line.fill.background()

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.28)
        ctf.margin_top = Inches(0.42)

        p = ctf.paragraphs[0]
        p.text = title
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_DARK
        p.font.name = "Arial"
        p.space_after = Pt(8)

        p_sub = ctf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.size = Pt(11.5)
        p_sub.font.color.rgb = C_TEXT_MUTED
        p_sub.font.name = "Arial"
        p_sub.space_after = Pt(14)

        for b in bullets:
            pb = ctf.add_paragraph()
            pb.text = "• " + b
            pb.font.size = Pt(11)
            pb.font.color.rgb = C_TEXT_DARK
            pb.font.name = "Arial"
            pb.space_after = Pt(7)

    # =========================================================================
    # SLIDE 3: THE SOLUTION (In-App Automation Testing)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3, C_BG_LIGHT)
    add_header(s3, "The Solution: In-App Automation Testing Platform", "OUR INNOVATION", 3)

    sol_box = s3.shapes.add_textbox(Inches(0.8), Inches(1.55), Inches(5.6), Inches(5.2))
    stf = sol_box.text_frame
    stf.word_wrap = True
    stf.margin_left = stf.margin_right = stf.margin_top = 0

    p = stf.paragraphs[0]
    p.text = "Bringing Test Automation Directly to Developers"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_DARK
    p.font.name = "Arial"
    p.space_after = Pt(8)

    p2 = stf.add_paragraph()
    p2.text = "Rather than pushing test execution off to remote CI queues, ocr-processor embeds the test harness into the running application under the dedicated Tests tab."
    p2.font.size = Pt(12)
    p2.font.color.rgb = C_TEXT_MUTED
    p2.font.name = "Arial"
    p2.space_after = Pt(16)

    sol_pillars = [
        ("Self-Provisioned Execution", "Test cases automatically seed isolated database records, mock external APIs, and establish authenticated sessions on demand."),
        ("Zero-Friction Selectivity", "Developers tick specific test cases or execute complete suites with an asynchronous non-blocking queue."),
        ("Triple-Fold Evidence", "Every run produces an audit breakdown (Executed / Validated / Observed) paired with real browser screenshots."),
        ("Requirements-Grounded Coverage", "Direct visual correlation between test runs and Phase 1 functional requirements (REQ-001 through REQ-007).")
    ]

    for title, desc in sol_pillars:
        pt = stf.add_paragraph()
        pt.text = "✔ " + title
        pt.font.size = Pt(12.5)
        pt.font.bold = True
        pt.font.color.rgb = C_BLUE
        pt.font.name = "Arial"

        pd = stf.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(11)
        pd.font.color.rgb = C_TEXT_DARK
        pd.font.name = "Arial"
        pd.space_after = Pt(8)

    # Right Column: Screenshot of the System Architecture / Product View
    img_arch = '/var/folders/fw/03lp5kq93p35tq26b7wxzmcc0000gn/T/cursor/screenshots/system-architecture.png'
    if os.path.exists(img_arch):
        pic_x = Inches(6.7)
        pic_y = Inches(1.6)
        pic_w = Inches(5.8)
        pic = s3.shapes.add_picture(img_arch, pic_x, pic_y, width=pic_w)
        
        cap = s3.shapes.add_textbox(pic_x, pic_y + pic.height + Inches(0.1), pic_w, Inches(0.35))
        cp = cap.text_frame.paragraphs[0]
        cp.text = "Figure 1: Application Architecture & End-to-End Testing Lifecycle"
        cp.font.size = Pt(10)
        cp.font.italic = True
        cp.font.color.rgb = C_TEXT_MUTED
        cp.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 4: THE PRODUCT UNDER TEST (Full Stack App)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4, C_BG_LIGHT)
    add_header(s4, "The Product Under Test: Document Intelligence Pipeline", "SYSTEM ARCHITECTURE", 4)

    features = [
        ("Multi-Format Ingestion", "Upload PDF, PNG, JPG, JPEG, and TIFF documents up to 20MB with real-time UI validation and progress states.", C_BLUE),
        ("Tesseract OCR Engine", "Local optical character recognition with ImageMagick preprocessing, DPI scaling, and TSV confidence filtering.", C_GREEN),
        ("PostgreSQL TSVECTOR Search", "High-speed full-text search indexed via GIN triggers across document names, OCR text, and summaries.", C_AMBER),
        ("Spring AI LLM Summaries", "Automatic document summarisation with smart token truncation (10,000 chars) and one-click retry for failed inferences.", C_BLUE),
        ("Hardened Security Layer", "Stateless authentication via HTTP-only JWT cookies, BCrypt password hashing (factor 12), and automatic 5-attempt account lockout.", C_RED),
        ("Interactive Web Interface", "Lightweight Thymeleaf + CSS grid frontend featuring live Search, Create, Profile, and Test Console views.", C_NAVY)
    ]

    for idx, (ftitle, fdesc, fcol) in enumerate(features):
        row = idx // 3
        col = idx % 3
        fx = Inches(0.8 + col * 4.0)
        fy = Inches(1.6 + row * 2.6)

        fcard = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, fx, fy, Inches(3.75), Inches(2.35))
        fcard.fill.solid()
        fcard.fill.fore_color.rgb = C_CARD_BG
        fcard.line.color.rgb = C_CARD_BORDER
        fcard.line.width = Pt(1.5)

        ftag = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, fx + Inches(0.2), fy + Inches(0.2), Inches(0.4), Inches(0.1))
        ftag.fill.solid()
        ftag.fill.fore_color.rgb = fcol
        ftag.line.fill.background()

        ftf = fcard.text_frame
        ftf.word_wrap = True
        ftf.margin_left = ftf.margin_right = Inches(0.25)
        ftf.margin_top = Inches(0.38)

        fp1 = ftf.paragraphs[0]
        fp1.text = ftitle
        fp1.font.size = Pt(14.5)
        fp1.font.bold = True
        fp1.font.color.rgb = C_TEXT_DARK
        fp1.space_after = Pt(6)

        fp2 = ftf.add_paragraph()
        fp2.text = fdesc
        fp2.font.size = Pt(11)
        fp2.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 5: TEST CONSOLE ARCHITECTURE (101 Cases)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5, C_BG_LIGHT)
    add_header(s5, "The Test Console: 101 Test Cases Across 3 Tiers", "AUTOMATED TEST TIERS", 5)

    tbox_s5 = s5.shapes.add_textbox(Inches(0.8), Inches(1.55), Inches(5.1), Inches(5.2))
    ttf = tbox_s5.text_frame
    ttf.word_wrap = True
    ttf.margin_left = ttf.margin_right = ttf.margin_top = 0

    tiers = [
        ("Unit Tests (32 Cases)", "Fast service-layer verification with faked dependencies. Validates JWT parsing & tamper rejection, lockout policies, truncation bounds, and OCR text filters.", C_BLUE),
        ("Regression Tests (29 Cases)", "Integration tests running against a real Testcontainers PostgreSQL instance. Proves full HTTP lifecycles, user isolation, cookie auth, and repository paging.", C_GREEN),
        ("End-to-End Tests (40 Cases)", "Real browser journeys driven by Selenium against the running application, combined with comprehensive API pytest scenarios testing user journeys.", C_AMBER)
    ]

    for tname, tdesc, tcolor in tiers:
        tp = ttf.add_paragraph() if ttf.paragraphs[0].text else ttf.paragraphs[0]
        tp.text = tname
        tp.font.size = Pt(14)
        tp.font.bold = True
        tp.font.color.rgb = tcolor
        tp.font.name = "Arial"
        tp.space_after = Pt(3)

        tpd = ttf.add_paragraph()
        tpd.text = tdesc
        tpd.font.size = Pt(11)
        tpd.font.color.rgb = C_TEXT_DARK
        tpd.font.name = "Arial"
        tpd.space_after = Pt(11)

    tpf = ttf.add_paragraph()
    tpf.text = "Console Capabilities:"
    tpf.font.size = Pt(12.5)
    tpf.font.bold = True
    tpf.font.color.rgb = C_TEXT_DARK
    tpf.space_after = Pt(3)

    caps = [
        "Selective checkbox picking or full suite execution",
        "Non-blocking background queue for simultaneous runs",
        "High-density two-row layout: zero horizontal scroll",
        "Failed cases automatically bubble to the top of table"
    ]
    for c in caps:
        cp = ttf.add_paragraph()
        cp.text = "• " + c
        cp.font.size = Pt(10.5)
        cp.font.color.rgb = C_TEXT_MUTED
        cp.space_after = Pt(2)

    img_console = '/var/folders/fw/03lp5kq93p35tq26b7wxzmcc0000gn/T/cursor/screenshots/two-row-failed-first.png'
    if os.path.exists(img_console):
        pic5_x = Inches(6.2)
        pic5_y = Inches(1.6)
        pic5_w = Inches(6.3)
        pic5 = s5.shapes.add_picture(img_console, pic5_x, pic5_y, width=pic5_w)

        cap5 = s5.shapes.add_textbox(pic5_x, pic5_y + pic5.height + Inches(0.1), pic5_w, Inches(0.35))
        cp5 = cap5.text_frame.paragraphs[0]
        cp5.text = "Figure 2: In-App Test Console showing Two-Row Case Layout & Failed-First Sorting"
        cp5.font.size = Pt(10)
        cp5.font.italic = True
        cp5.font.color.rgb = C_TEXT_MUTED
        cp5.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 6: ACTIONABLE EVIDENCE & DIAGNOSTICS
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6, C_BG_LIGHT)
    add_header(s6, "Actionable Evidence & Root-Cause Failure Diagnostics", "EVIDENCE & DIAGNOSTICS", 6)

    ev_cards = [
        ("EXECUTED", "Input Parameters", "Records user ID, uploaded filename, MIME type, auth headers, and execution flags.", C_BLUE),
        ("VALIDATED", "Assertions & Contracts", "Documents required HTTP status, database persistence, and contract preconditions.", C_GREEN),
        ("OBSERVED", "Runtime Outputs", "Captures server response JSON, OCR text excerpts, and error messages.", C_AMBER)
    ]

    for i, (etag, etitle, edesc, ecol) in enumerate(ev_cards):
        ex = Inches(0.8 + i * 1.95)
        ecard = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, ex, Inches(1.6), Inches(1.8), Inches(2.2))
        ecard.fill.solid()
        ecard.fill.fore_color.rgb = C_CARD_BG
        ecard.line.color.rgb = C_CARD_BORDER
        ecard.line.width = Pt(1.5)

        etf = ecard.text_frame
        etf.word_wrap = True
        etf.margin_left = etf.margin_right = Inches(0.15)
        etf.margin_top = Inches(0.18)

        ep1 = etf.paragraphs[0]
        ep1.text = etag
        ep1.font.size = Pt(10.5)
        ep1.font.bold = True
        ep1.font.color.rgb = ecol

        ep2 = etf.add_paragraph()
        ep2.text = etitle
        ep2.font.size = Pt(11.5)
        ep2.font.bold = True
        ep2.font.color.rgb = C_TEXT_DARK
        ep2.space_after = Pt(4)

        ep3 = etf.add_paragraph()
        ep3.text = edesc
        ep3.font.size = Pt(9.5)
        ep3.font.color.rgb = C_TEXT_MUTED

    fcard = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.0), Inches(5.7), Inches(2.8))
    fcard.fill.solid()
    fcard.fill.fore_color.rgb = C_RED_BG
    fcard.line.color.rgb = RGBColor(252, 165, 165)
    fcard.line.width = Pt(1.5)

    ftf = fcard.text_frame
    ftf.word_wrap = True
    ftf.margin_left = ftf.margin_right = Inches(0.25)
    ftf.margin_top = Inches(0.22)

    fp = ftf.paragraphs[0]
    fp.text = "Deliberate Failure Diagnostics (Where · Why · How to Fix)"
    fp.font.size = Pt(13.5)
    fp.font.bold = True
    fp.font.color.rgb = C_RED
    fp.space_after = Pt(6)

    diag_bullets = [
        "Where: pinpoints the exact file and line number in source code (e.g., FailureSamplesTest.java:17)",
        "Why: contrasts expected contract against observed reality (e.g., Expected length >= 8, but was 3)",
        "How to Fix: guides the developer on the exact architectural remediation required",
        "Evidence Card: every non-UI test auto-generates a 640px visual summary card for audits"
    ]
    for db in diag_bullets:
        dbp = ftf.add_paragraph()
        dbp.text = "• " + db
        dbp.font.size = Pt(10)
        dbp.font.color.rgb = C_TEXT_DARK
        dbp.space_after = Pt(4)

    img_fail = '/var/folders/fw/03lp5kq93p35tq26b7wxzmcc0000gn/T/cursor/screenshots/failure-details.png'
    if os.path.exists(img_fail):
        pic6_x = Inches(6.8)
        pic6_y = Inches(1.6)
        pic6_w = Inches(5.7)
        pic6 = s6.shapes.add_picture(img_fail, pic6_x, pic6_y, width=pic6_w)

        cap6 = s6.shapes.add_textbox(pic6_x, pic6_y + pic6.height + Inches(0.1), pic6_w, Inches(0.35))
        cp6 = cap6.text_frame.paragraphs[0]
        cp6.text = "Figure 3: Interactive Failure Explanation Breakdown & Stack Tracing"
        cp6.font.size = Pt(10)
        cp6.font.italic = True
        cp6.font.color.rgb = C_TEXT_MUTED
        cp6.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 7: REQUIREMENT-DRIVEN COVERAGE
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7, C_BG_LIGHT)
    add_header(s7, "Requirements-Phase Traceability: Beyond Raw Line Coverage", "COVERAGE METRICS", 7)

    cov_box = s7.shapes.add_textbox(Inches(0.8), Inches(1.55), Inches(5.2), Inches(5.2))
    ctf7 = cov_box.text_frame
    ctf7.word_wrap = True
    ctf7.margin_left = ctf7.margin_right = ctf7.margin_top = 0

    cp1 = ctf7.paragraphs[0]
    cp1.text = "38 of 52 Criteria Verified (73.1%)"
    cp1.font.size = Pt(17)
    cp1.font.bold = True
    cp1.font.color.rgb = C_GREEN
    cp1.space_after = Pt(6)

    cp2 = ctf7.add_paragraph()
    cp2.text = "Traditional code coverage tools report lines executed without proving requirements. The Coverage tab maps automated cases directly to acceptance criteria defined in Phase 1:"
    cp2.font.size = Pt(11.5)
    cp2.font.color.rgb = C_TEXT_MUTED
    cp2.space_after = Pt(10)

    req_summary = [
        ("REQ-001 Registration (66.7%)", "Duplicate checks, short passwords, redirects."),
        ("REQ-002 Login (62.5%)", "Valid auth, generic error messaging, protected routes."),
        ("REQ-003 Change Password (71.4%)", "Validation, current password verification."),
        ("REQ-004 Home UI (80.0%)", "Tabs, navigation, and seamless logout flows."),
        ("REQ-005 OCR Search (62.5%)", "Queries, empty states, and tenant isolation."),
        ("REQ-006 Upload & OCR (87.5%)", "Ingestion, text extraction, and status transitions."),
        ("REQ-007 LLM Summary (85.7%)", "Automatic summarisation and failure retries.")
    ]
    for rname, rdesc in req_summary:
        rp = ctf7.add_paragraph()
        rp.text = "• " + rname + ": " + rdesc
        rp.font.size = Pt(10)
        rp.font.color.rgb = C_TEXT_DARK
        rp.space_after = Pt(2.5)

    cp_open = ctf7.add_paragraph()
    cp_open.text = "\nHonest Accounting of Open Gaps:\nExplicitly explains why remaining criteria (e.g., >20MB uploads, live external LLM billing, Tesseract 2-min timeouts) remain open."
    cp_open.font.size = Pt(10)
    cp_open.font.italic = True
    cp_open.font.color.rgb = C_TEXT_MUTED

    img_req = '/var/folders/fw/03lp5kq93p35tq26b7wxzmcc0000gn/T/cursor/screenshots/req-coverage-top.png'
    if os.path.exists(img_req):
        pic7_x = Inches(6.3)
        pic7_y = Inches(1.6)
        pic7_w = Inches(6.2)
        pic7 = s7.shapes.add_picture(img_req, pic7_x, pic7_y, width=pic7_w)

        cap7 = s7.shapes.add_textbox(pic7_x, pic7_y + pic7.height + Inches(0.1), pic7_w, Inches(0.35))
        cp7 = cap7.text_frame.paragraphs[0]
        cp7.text = "Figure 4: Visual Breakdown of Requirements Coverage & Rationales for Open Items"
        cp7.font.size = Pt(10)
        cp7.font.italic = True
        cp7.font.color.rgb = C_TEXT_MUTED
        cp7.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 8: CO-PILOTED WITH IBM BOB
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8, C_BG_LIGHT)
    add_header(s8, "Co-Piloted with IBM Bob: From Concept to Working System", "IBM BOB METHODOLOGY", 8)

    phases = [
        ("Phase 1: Requirements Engineering",
         "Captured 7 user stories, EARS acceptance criteria, edge cases, and data models in structured documentation.",
         C_BLUE),
        ("Phase 2: Architecture & Decision Records",
         "Authored 5 Architectural Decision Records (ADRs) evaluating Tesseract, Spring AI, PostgreSQL FTS, and JWT cookies.",
         C_GREEN),
        ("Phase 3: Design & Technical Spec",
         "Generated end-to-end component wireframes, sequence diagrams, and REST contract schemas before writing code.",
         C_AMBER),
        ("Phase 4: Full-Stack Implementation",
         "Implemented Spring Boot application, Testcontainers infrastructure, Pytest Selenium runners, and the interactive Test Console.",
         C_NAVY)
    ]

    for idx, (pname, pdesc, pcol) in enumerate(phases):
        px = Inches(0.8 + idx * 2.95)
        pcard = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, Inches(1.7), Inches(2.75), Inches(4.9))
        pcard.fill.solid()
        pcard.fill.fore_color.rgb = C_CARD_BG
        pcard.line.color.rgb = C_CARD_BORDER
        pcard.line.width = Pt(1.5)

        pnum = s8.shapes.add_shape(MSO_SHAPE.OVAL, px + Inches(0.25), Inches(1.95), Inches(0.6), Inches(0.6))
        pnum.fill.solid()
        pnum.fill.fore_color.rgb = pcol
        pnum.line.fill.background()
        pn_tf = pnum.text_frame
        pn_p = pn_tf.paragraphs[0]
        pn_p.text = str(idx + 1)
        pn_p.font.size = Pt(14)
        pn_p.font.bold = True
        pn_p.font.color.rgb = C_WHITE
        pn_p.alignment = PP_ALIGN.CENTER

        ptf8 = pcard.text_frame
        ptf8.word_wrap = True
        ptf8.margin_left = ptf8.margin_right = Inches(0.25)
        ptf8.margin_top = Inches(0.95)

        p1 = ptf8.paragraphs[0]
        p1.text = pname
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = C_TEXT_DARK
        p1.space_after = Pt(8)

        p2 = ptf8.add_paragraph()
        p2.text = pdesc
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_TEXT_MUTED
        p2.space_after = Pt(12)

        p3 = ptf8.add_paragraph()
        p3.text = "Key Deliverable: Markdown Spec & ADRs" if idx < 3 else "Key Deliverable: 101 Executable Cases"
        p3.font.size = Pt(10)
        p3.font.bold = True
        p3.font.color.rgb = pcol

    # =========================================================================
    # SLIDE 9: SDLC IMPACT & EFFICIENCY
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_bg(s9, C_BG_LIGHT)
    add_header(s9, "Measurable SDLC Impact: Accelerating Release Velocity", "BUSINESS & TECHNICAL IMPACT", 9)

    impact_metrics = [
        ("Feedback Latency", "Seconds vs. Days", "Developers receive immediate failure attribution in their local browser session instead of waiting for nightlies or QA ticket cycles.", C_GREEN),
        ("Regression Confidence", "Zero Silent Breaks", "Executing the 29-case regression suite proves auth, search, upload, and database isolation before committing changes.", C_BLUE),
        ("Diagnostic Overhead", "Eliminated Guesswork", "Every failure includes source line coordinates, contract discrepancy reasons, and visual evidence cards.", C_AMBER),
        ("Handoff Quality", "Clean QA Drops", "Testing teams receive verified builds where baseline functionality is proven, freeing QA for exploratory testing.", C_NAVY)
    ]

    for idx, (mtitle, mstat, mdesc, mcol) in enumerate(impact_metrics):
        row = idx // 2
        col = idx % 2
        mx = Inches(0.8 + col * 6.0)
        my = Inches(1.7 + row * 2.5)

        mcard = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, mx, my, Inches(5.7), Inches(2.25))
        mcard.fill.solid()
        mcard.fill.fore_color.rgb = C_CARD_BG
        mcard.line.color.rgb = C_CARD_BORDER
        mcard.line.width = Pt(1.5)

        mtf = mcard.text_frame
        mtf.word_wrap = True
        mtf.margin_left = mtf.margin_right = Inches(0.3)
        mtf.margin_top = Inches(0.28)

        mp1 = mtf.paragraphs[0]
        mp1.text = mtitle.upper()
        mp1.font.size = Pt(11)
        mp1.font.bold = True
        mp1.font.color.rgb = mcol

        mp2 = mtf.add_paragraph()
        mp2.text = mstat
        mp2.font.size = Pt(18)
        mp2.font.bold = True
        mp2.font.color.rgb = C_TEXT_DARK
        mp2.space_after = Pt(5)

        mp3 = mtf.add_paragraph()
        mp3.text = mdesc
        mp3.font.size = Pt(11)
        mp3.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 10: CONCLUSION & SUMMARY (Dark Navy)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_bg(s10, C_NAVY)

    top_line10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.15), Inches(1.5), Inches(0.08))
    top_line10.fill.solid()
    top_line10.fill.fore_color.rgb = C_GREEN
    top_line10.line.fill.background()

    cbox = s10.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.7), Inches(4.8))
    ctf10 = cbox.text_frame
    ctf10.word_wrap = True

    cp1 = ctf10.paragraphs[0]
    cp1.text = "Summary: Testing as an Active Development Tool"
    cp1.font.size = Pt(36)
    cp1.font.bold = True
    cp1.font.color.rgb = C_WHITE
    cp1.space_after = Pt(18)

    cp2 = ctf10.add_paragraph()
    cp2.text = "• Complete Transparency: Every test is self-documenting with Executed, Validated, and Observed evidence.\n" \
               "• Integrated Developer Experience: Run Unit, Regression, and End-to-End suites without leaving the application.\n" \
               "• Requirements-Driven Quality: Continuous tracking against formal acceptance criteria (REQ-001..REQ-007).\n" \
               "• Accelerated SDLC: Catch defects earlier, eliminate handoff friction, and deliver faster with confidence."
    cp2.font.size = Pt(16)
    cp2.font.color.rgb = RGBColor(226, 232, 240)
    cp2.space_after = Pt(28)

    cp3 = ctf10.add_paragraph()
    cp3.text = "Thank you! Ready for Live Demo & Q&A"
    cp3.font.size = Pt(22)
    cp3.font.bold = True
    cp3.font.color.rgb = RGBColor(56, 189, 248)

    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "ocr-processor-presentation.pptx"
    create_deck(out)
